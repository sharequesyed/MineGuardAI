import os
import sys

try:
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT

DOCX_PATH = os.path.join(os.path.dirname(__file__), "MineGuard-AI_Hardware_Manual.docx")

def build_hardware_manual():
    doc = docx.Document()

    # Define Styles & Colors
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    # Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_p.add_run("MINEGUARD-AI HARDWARE ASSEMBLY & SPECIFICATION MANUAL")
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = subtitle_p.add_run("Smart India Hackathon 2026 — Problem Statement SIH26025\nStudent Prototype Platform Specification & Construction Guide")
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    doc.add_paragraph() # Spacer

    # Notice Alert Box
    p_notice = doc.add_paragraph()
    r_notice = p_notice.add_run("STUDENT PROTOTYPE NOTICE:\nMineGuard-AI is a student prototype platform designed toward a scalable field architecture. It does NOT possess mine certification, intrinsic safety ratings, environmental qualification, or field-validated collapse prediction accuracy. All pin mappings are marked PROVISIONAL pending exact board choice.")
    r_notice.font.bold = True
    r_notice.font.size = Pt(10)
    r_notice.font.color.rgb = RGBColor(0xB4, 0x53, 0x09)

    doc.add_heading("1. Project Hardware Overview", level=1)
    doc.add_paragraph("MineGuard-AI addresses surface subsidence caused by underground coal mining. 4 prototype sensor nodes (N1–N4) and 1 ESP32 gateway are deployed ON THE SURFACE ABOVE underground coal panels to measure surface inclination, relative displacement delta between surface points, surface vibration, and continuity line crack initiation.")

    doc.add_heading("2. Safety & Electrical Notes", level=1)
    doc.add_paragraph("Always disconnect USB power before modifying wiring. Verify 3.3V vs 5V pin operating levels on the ESP32. Do NOT apply 5V directly to ESP32 GPIO input pins.")

    doc.add_heading("3. Complete Bill of Materials (BOM)", level=1)
    
    # Table BOM
    table_bom = doc.add_table(rows=1, cols=5)
    table_bom.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table_bom.rows[0].cells
    headers = ["Item #", "Component", "Quantity", "Est. Unit Cost (INR)", "Purpose"]
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        hdr_cells[i].paragraphs[0].runs[0].font.bold = True

    bom_data = [
        ["1", "ESP32 Development Board (30-pin WROOM-32)", "5 (4 nodes + 1 GW)", "₹350", "Microcontroller & Wireless ESP-NOW Communication"],
        ["2", "MPU6050 6-DOF Accelerometer & Gyroscope", "4", "₹120", "Surface tilt inclination & vibration RMS sensing"],
        ["3", "Linear Slider / String Potentiometer (10k)", "4", "₹80", "Relative displacement delta between surface node pairs"],
        ["4", "Conductive Continuity Line / Break Wire", "4", "₹15", "Prototype crack detector continuity bridge"],
        ["5", "Active Piezo Buzzer & Indicator LED", "4", "₹25", "Local fail-safe threshold alarm trigger"],
        ["6", "Solderless Breadboards & Jumper Wires", "1 Set", "₹200", "Prototype circuit assembly"],
        ["7", "USB Micro/Type-C Cables", "2", "₹100", "Gateway USB Web Serial power and data connection"],
    ]

    for row in bom_data:
        row_cells = table_bom.add_row().cells
        for i, val in enumerate(row):
            row_cells[i].text = val

    doc.add_heading("4. Exact Pin Mapping (PROVISIONAL / VERIFY WITH ACTUAL BOARD)", level=1)
    doc.add_paragraph("Pin definitions are centralized in hardware_config.h. The table below outlines the provisional pin configuration for ESP32-WROOM-32.")

    table_pin = doc.add_table(rows=1, cols=4)
    table_pin.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_pin = table_pin.rows[0].cells
    pin_headers = ["Component", "Component Pin", "ESP32 GPIO", "Electrical Purpose / Constraints"]
    for i, h in enumerate(pin_headers):
        hdr_pin[i].text = h
        hdr_pin[i].paragraphs[0].runs[0].font.bold = True

    pin_data = [
        ["MPU6050", "SDA", "GPIO 21", "I2C Data line (3.3V)"],
        ["MPU6050", "SCL", "GPIO 22", "I2C Clock line (3.3V)"],
        ["String Potentiometer", "Wiper Out", "GPIO 34 (ADC1_CH6)", "Analog displacement input (Input-Only, 0-3.3V)"],
        ["Crack Detector", "Continuity Return", "GPIO 4", "Digital Input with Internal Pull-up (LOW=Intact, HIGH=Broken)"],
        ["Local Buzzer", "VCC Positive", "GPIO 18", "Digital Output (High level 3.3V trigger)"],
        ["ESP-NOW Gateway", "USB Serial RX/TX", "UART 0 (USB)", "Web Serial JSON telemetry output to laptop"],
    ]

    for row in pin_data:
        row_cells = table_pin.add_row().cells
        for i, val in enumerate(row):
            row_cells[i].text = val

    doc.add_heading("5. Physical Placement & Surface Model Construction", level=1)
    doc.add_paragraph("Sensors are mounted on a controlled surface ground model above an underground coal extraction panel. The 4 nodes form a rectangle: Node N1 (top-left), Node N2 (top-right), Node N3 (bottom-right), Node N4 (bottom-left). String potentiometers span between N1-N2, N2-N3, N3-N4, N1-N4 to measure relative ground movement.")

    doc.add_heading("6. Sensor Calibration Procedure", level=1)
    doc.add_paragraph("1. MPU6050 Level Zeroing: Place node on horizontal surface; record X/Y zero offset.\n2. Displacement Mechanism Zeroing: Extend string potentiometer to neutral baseline; set baseline ADC reference.\n3. Crack Detector Test: Verify digital input pin reads LOW when conductive wire bridge is intact.")

    doc.add_heading("7. Step-by-Step SIH Demonstration Procedure", level=1)
    doc.add_paragraph("1. Connect ESP32 Gateway to laptop over USB cable.\n2. Open MineGuard-AI web application in Chrome browser.\n3. Click [Connect Hardware] -> Web Serial -> Select ESP32 Port.\n4. Observe 4 surface nodes NORMAL baseline.\n5. Apply surface tilt to Node N3 -> Web app displays WARNING.\n6. Pull displacement mechanism between N3 and N4 -> Web app displays CRITICAL & Turf.js spatial risk boundary polygon.\n7. Disconnect conductive wire on N4 -> Local buzzer triggers & web alert displays CRACK EVENT.\n8. Toggle offline mode -> System stores readings to IndexedDB queue and syncs when reconnected.")

    doc.add_heading("8. Troubleshooting Guide", level=1)
    doc.add_paragraph("Issue: ESP32 Port Not Detected -> Ensure CP210x / CH340 USB driver is installed.\nIssue: Web Serial Permission Denied -> Grant browser serial access when prompted.\nIssue: MPU6050 Read Error -> Verify SDA/SCL pullup resistors and 3.3V power connection.")

    doc.save(DOCX_PATH)
    print(f"MineGuard-AI Hardware Manual DOCX created successfully at {DOCX_PATH}")

if __name__ == "__main__":
    build_hardware_manual()
